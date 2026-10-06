"""Load a reviewed translation sheet back into app/scripts/i18n/strings.js.

usage: python tools/i18n_import.py reviewed.xlsx [--dry-run]

Matches rows by the English text. Checks that every {placeholder} in the
English text is still present in the Korean and Arabic text, and reports
rows whose English no longer exists in the code. Prints a summary of what
changed; with --dry-run nothing is written.
"""
import io, json, os, re, sys
from openpyxl import load_workbook

HERE = os.path.dirname(__file__)
STRINGS = os.path.join(HERE, '..', 'app', 'scripts', 'i18n', 'strings.js')
PH = re.compile(r'\{\w+\}')

def main():
    path = sys.argv[1]
    dry = '--dry-run' in sys.argv
    src = io.open(STRINGS, encoding='utf-8').read()
    head = src[:src.index('export default')]
    data = json.loads(src[src.index('export default') + len('export default'):].strip().rstrip(';'))
    ws = load_workbook(path)['Translations']
    changed, problems, unknown = [], [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue
        en, ko, ar = str(row[1]), (row[2] or '').strip(), (row[3] or '').strip()
        if en not in data:
            unknown.append(en); continue
        need = set(PH.findall(en))
        for lang, txt in (('ko', ko), ('ar', ar)):
            if txt and set(PH.findall(txt)) != need:
                problems.append(f'{lang}: placeholders differ — "{en}" -> "{txt}"')
                continue
            if txt and data[en].get(lang) != txt:
                changed.append(f'{lang}: "{en}"  {data[en].get(lang)!r} -> {txt!r}')
                data[en][lang] = txt
    print(f'{len(changed)} changed · {len(problems)} rejected (placeholders) · {len(unknown)} not in the code')
    for x in changed: print('  ~', x)
    for x in problems: print('  !', x)
    for x in unknown: print('  ?', x)
    if not dry and changed:
        io.open(STRINGS, 'w', encoding='utf-8', newline='\n').write(head + 'export default ' + json.dumps(data, ensure_ascii=False, indent=1) + ';\n')
        print('strings.js updated')

main()
