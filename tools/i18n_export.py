"""Export the interface translations to an Excel review sheet.

usage: python tools/i18n_export.py [out.xlsx]
default out: docs/translations/PPMS_UI_translations.xlsx

Reviewers edit only the Korean / Arabic columns (and Notes); English is the
key and must not be changed. Send the file back and run tools/i18n_import.py.
"""
import io, json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(__file__)
STRINGS = os.path.join(HERE, '..', 'app', 'scripts', 'i18n', 'strings.js')
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, '..', 'docs', 'translations', 'PPMS_UI_translations.xlsx')

GLOSSARY = [
    ('block (a bar on the Gantt)', '블록', 'بلوك'),
    ('process', '공정', 'عملية'),
    ('station', '스테이션', 'محطة'),
    ('vehicle', '차량', 'مركبة'),
    ('unit (e.g. M1)', '유닛', 'وحدة'),
    ('battalion', '대대', 'كتيبة'),
    ('Hull / Turret / Assembly', '차체 / 포탑 / 조립', 'الهيكل / البرج / التجميع'),
    ('component (Hull, Turret, Assembly)', '구성품', 'مكوّن'),
    ('line (production line)', '라인', 'خط'),
    ('plan version', '계획 버전', 'إصدار الخطة'),
    ('overdue', '기한 초과', 'متأخر عن الموعد'),
    ('completed late / early', '지연 완료 / 조기 완료', 'اكتمل متأخرًا / مبكرًا'),
    ('working days (wd)', '근무일', 'يوم عمل'),
    ('reschedule', '일정 변경', 'إعادة الجدولة'),
    ('Executive Summary / Report', '경영 요약 / 보고서', 'الملخص / التقرير التنفيذي'),
    ('Gantt', '간트', 'جانت'),
    ('VPX, K9, K10, K11, BTL-01 (codes)', 'do not translate', 'لا تُترجم'),
]

def load():
    src = io.open(STRINGS, encoding='utf-8').read()
    return json.loads(src[src.index('export default') + len('export default'):].strip().rstrip(';'))

def main():
    data = load()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Translations'
    head = ['Area', 'English (do not edit)', 'Korean 한국어', 'Arabic العربية', 'Reviewer notes']
    ws.append(head)
    hfill = PatternFill('solid', fgColor='1E293B')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = hfill
        c.alignment = Alignment(vertical='center')
    thin = Side(style='thin', color='E2E8F0')
    area_fill = {'Header & menus': 'EFF6FF', 'Filter bar': 'F0FDF4', 'Executive Summary': 'FEFCE8',
                 'Gantt': 'FDF4FF', 'Gantt / shared': 'FAF5FF', 'Sign-in page': 'F1F5F9'}
    for text, e in data.items():
        ws.append([e.get('area', ''), text, e.get('ko', ''), e.get('ar', ''), ''])
        r = ws.max_row
        for col in range(1, 6):
            c = ws.cell(r, col)
            c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='right' if col == 4 else 'left',
                                    readingOrder=2 if col == 4 else 0)
            c.border = Border(bottom=thin)
        ws.cell(r, 1).fill = PatternFill('solid', fgColor=area_fill.get(e.get('area'), 'FFFFFF'))
        ws.cell(r, 2).font = Font(color='475569')
    for col, w in zip('ABCDE', (18, 60, 50, 50, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:E{ws.max_row}'

    g = wb.create_sheet('Glossary')
    g.append(['English term', 'Korean', 'Arabic'])
    for c in g[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = hfill
    for row in GLOSSARY:
        g.append(list(row))
        g.cell(g.max_row, 3).alignment = Alignment(horizontal='right', readingOrder=2)
    for col, w in zip('ABC', (40, 30, 30)):
        g.column_dimensions[col].width = w

    h = wb.create_sheet('How to review')
    lines = [
        'PPMS interface translations — review guide',
        '',
        '1. Edit only the Korean and Arabic columns (and Reviewer notes). Do not change the English column — it is the key.',
        '2. Keep placeholders exactly as written: {n}, {b}, {a}, {done}, {total}. They are replaced by numbers or words on screen.',
        '3. Keep codes and names untranslated: VPX, K9, K10, K11, BTL-01, M1, station and part names.',
        '4. Use the Glossary sheet so the same term is always translated the same way.',
        '5. Short is better: most texts are buttons and labels with little space.',
        '6. Arabic: Western digits (0-9); the Gantt and VPX stay left-to-right, everything else is right-to-left.',
        '7. If a text is unclear, write a question in Reviewer notes and leave the translation as is.',
        '8. Send the file back; it is loaded with tools/i18n_import.py.',
    ]
    for l in lines:
        h.append([l])
    h['A1'].font = Font(bold=True, size=13)
    h.column_dimensions['A'].width = 120

    os.makedirs(os.path.dirname(os.path.abspath(OUT)), exist_ok=True)
    wb.save(OUT)
    print(f'{len(data)} texts -> {os.path.abspath(OUT)}')

main()
